import re
import markdown
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt, Emu
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import PP_ALIGN
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.units import mm
import html


class MarkdownToDocument:
    def to_docx(self, md_text: str, output_path: str):
        doc = Document()
        # Set default font
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Microsoft YaHei'
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        lines = md_text.strip().split('\n')
        i = 0
        while i < len(lines):
            line = lines[i]

            # Headings
            if line.startswith('# ') and not line.startswith('## '):
                p = doc.add_heading(line[2:].strip(), level=1)
            elif line.startswith('## ') and not line.startswith('### '):
                p = doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith('### '):
                p = doc.add_heading(line[4:].strip(), level=3)
            # Unordered list
            elif re.match(r'^[\-\*]\s', line):
                doc.add_paragraph(line.lstrip('-* ').strip(), style='List Bullet')
            # Ordered list
            elif re.match(r'^\d+\.\s', line):
                doc.add_paragraph(re.sub(r'^\d+\.\s*', '', line), style='List Number')
            # Table
            elif line.strip().startswith('|') and i + 1 < len(lines) and lines[i + 1].strip().startswith('|-'):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    row_text = lines[i].strip().strip('|')
                    cells = [c.strip() for c in row_text.split('|')]
                    if not all(c.replace('-', '').replace(':', '').strip() == '' for c in cells):
                        table_lines.append(cells)
                    i += 1
                if table_lines:
                    cols = len(table_lines[0])
                    table = doc.add_table(rows=len(table_lines), cols=cols)
                    table.style = 'Light Grid Accent 1'
                    for r, row_cells in enumerate(table_lines):
                        for c, cell_text in enumerate(row_cells):
                            if c < cols:
                                table.rows[r].cells[c].text = cell_text
                    i -= 1
            # Code block
            elif line.strip().startswith('```'):
                i += 1
                code_lines = []
                while i < len(lines) and not lines[i].strip().startswith('```'):
                    code_lines.append(lines[i])
                    i += 1
                doc.add_paragraph('\n'.join(code_lines), style='No Spacing')
            # Empty line
            elif line.strip() == '':
                pass
            # Normal paragraph with inline formatting
            else:
                p = doc.add_paragraph()
                p.text = self._strip_inline_formatting(line.strip())

            i += 1

        doc.save(output_path)

    def to_pptx(self, md_text: str, output_path: str):
        prs = Presentation()
        prs.slide_width = PptxInches(13.333)
        prs.slide_height = PptxInches(7.5)

        blank_layout = prs.slide_layouts[6]  # Blank layout
        lines = md_text.strip().split('\n')

        slide = None

        for line in lines:
            line_stripped = line.strip()

            if not line_stripped:
                continue

            # Slide break: H1 creates a new slide
            if line_stripped.startswith('# ') and not line_stripped.startswith('## '):
                slide = prs.slides.add_slide(blank_layout)
                tf = slide.shapes.add_textbox(PptxInches(0.5), PptxInches(0.3), PptxInches(12), PptxInches(1))
                p = tf.text_frame.paragraphs[0]
                p.text = line_stripped[2:].strip()
                p.font.size = PptxPt(32)
                p.font.bold = True
                p.font.color.rgb = PptxRGBColor(0x33, 0x33, 0x33)

            elif line_stripped.startswith('## '):
                if slide is None:
                    slide = prs.slides.add_slide(blank_layout)
                tf = slide.shapes.add_textbox(PptxInches(0.5), PptxInches(0.3), PptxInches(12), PptxInches(0.8))
                p = tf.text_frame.paragraphs[0]
                p.text = line_stripped[3:].strip()
                p.font.size = PptxPt(24)
                p.font.bold = True

            elif line_stripped.startswith('### '):
                if slide is None:
                    slide = prs.slides.add_slide(blank_layout)
                tf = slide.shapes.add_textbox(PptxInches(0.5), PptxInches(1.0), PptxInches(12), PptxInches(0.6))
                p = tf.text_frame.paragraphs[0]
                p.text = line_stripped[4:].strip()
                p.font.size = PptxPt(20)
                p.font.bold = True

            else:
                if slide is None:
                    slide = prs.slides.add_slide(blank_layout)
                text = self._strip_inline_formatting(line_stripped)
                if re.match(r'^[\-\*]\s', line_stripped) or re.match(r'^\d+\.\s', line_stripped):
                    text = '  • ' + text.lstrip('-*0123456789. ')
                # Check if there's already a textbox at the bottom
                shapes_with_text = [s for s in slide.shapes if s.has_text_frame]
                if shapes_with_text:
                    last_shape = shapes_with_text[-1]
                    p = last_shape.text_frame.add_paragraph()
                    p.text = text
                    p.font.size = PptxPt(16)
                    p.font.color.rgb = PptxRGBColor(0x55, 0x55, 0x55)
                else:
                    tf = slide.shapes.add_textbox(PptxInches(0.5), PptxInches(1.5), PptxInches(12), PptxInches(5))
                    p = tf.text_frame.paragraphs[0]
                    p.text = text
                    p.font.size = PptxPt(16)
                    p.font.color.rgb = PptxRGBColor(0x55, 0x55, 0x55)

        if len(prs.slides) == 0:
            slide = prs.slides.add_slide(blank_layout)

        prs.save(output_path)

    def to_html(self, md_text: str, output_path: str):
        html_body = markdown.markdown(md_text, extensions=['tables', 'fenced_code', 'codehilite'])
        full_html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Converted Document</title>
<style>
body {{
    max-width: 800px;
    margin: 40px auto;
    padding: 0 20px;
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
    color: #333;
    line-height: 1.6;
}}
h1 {{ color: #1a1a1a; border-bottom: 2px solid #6366f1; padding-bottom: 8px; }}
h2 {{ color: #333; border-bottom: 1px solid #ddd; padding-bottom: 6px; }}
table {{ border-collapse: collapse; width: 100%; margin: 16px 0; }}
th, td {{ border: 1px solid #ddd; padding: 8px 12px; text-align: left; }}
th {{ background: #f5f5f5; font-weight: 600; }}
tr:nth-child(even) {{ background: #fafafa; }}
code {{ background: #f4f4f4; padding: 2px 6px; border-radius: 3px; font-size: 0.9em; }}
pre {{ background: #f4f4f4; padding: 16px; border-radius: 6px; overflow-x: auto; }}
pre code {{ background: none; padding: 0; }}
</style>
</head>
<body>
{html_body}
</body>
</html>"""
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(full_html)

    def to_pdf(self, md_text: str, output_path: str):
        doc = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            leftMargin=25 * mm,
            rightMargin=25 * mm,
            topMargin=20 * mm,
            bottomMargin=20 * mm,
        )

        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(
            name='CustomH1',
            parent=styles['Heading1'],
            fontSize=20,
            spaceAfter=12,
            textColor=HexColor('#1a1a1a'),
        ))
        styles.add(ParagraphStyle(
            name='CustomH2',
            parent=styles['Heading2'],
            fontSize=16,
            spaceAfter=8,
            textColor=HexColor('#333333'),
        ))
        styles.add(ParagraphStyle(
            name='CustomBody',
            parent=styles['Normal'],
            fontSize=11,
            spaceAfter=6,
            leading=16,
            textColor=HexColor('#333333'),
        ))
        styles.add(ParagraphStyle(
            name='CustomCode',
            parent=styles['Normal'],
            fontName='Courier',
            fontSize=9,
            backColor=HexColor('#f4f4f4'),
            spaceAfter=6,
            leading=13,
        ))

        html_body = markdown.markdown(md_text, extensions=['tables', 'fenced_code'])
        story = []

        for line in html_body.split('\n'):
            stripped = line.strip()
            if not stripped:
                story.append(Spacer(1, 6))
                continue

            # Headings
            if stripped.startswith('<h1>'):
                text = html.unescape(re.sub(r'</?h1>', '', stripped))
                story.append(Paragraph(text, styles['CustomH1']))
            elif stripped.startswith('<h2>'):
                text = html.unescape(re.sub(r'</?h2>', '', stripped))
                story.append(Paragraph(text, styles['CustomH2']))
            elif stripped.startswith('<h3>'):
                text = html.unescape(re.sub(r'</?h3>', '', stripped))
                story.append(Paragraph(text, styles['Heading3']))
            # Code blocks
            elif stripped.startswith('<pre>') or stripped.startswith('<code'):
                code_text = re.sub(r'<[^>]+>', '', stripped)
                story.append(Paragraph(code_text, styles['CustomCode']))
            # List items
            elif stripped.startswith('<li>'):
                text = html.unescape(re.sub(r'<[^>]+>', '', stripped))
                story.append(Paragraph(f'  • {text}', styles['CustomBody']))
            # Table rows
            elif stripped.startswith('<tr'):
                cells = re.findall(r'<t[dh][^>]*>(.*?)</t[dh]>', stripped)
                text = '  |  '.join(html.unescape(c) for c in cells)
                story.append(Paragraph(text, styles['CustomBody']))
            # Regular paragraph
            elif stripped.startswith('<p>'):
                text = html.unescape(re.sub(r'<[^>]+>', '', stripped))
                story.append(Paragraph(text, styles['CustomBody']))
            else:
                text = html.unescape(re.sub(r'<[^>]+>', '', stripped))
                if text:
                    story.append(Paragraph(text, styles['CustomBody']))

        doc.build(story)

    @staticmethod
    def _strip_inline_formatting(text: str) -> str:
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        text = re.sub(r'\*(.+?)\*', r'\1', text)
        text = re.sub(r'`(.+?)`', r'\1', text)
        return text
